import openpyxl
import os
import sys
import json
import re

# Ensure standard output uses UTF-8 encoding for Devanagari Unicode
sys.stdout.reconfigure(encoding='utf-8')

if len(sys.argv) < 2:
    print(json.dumps({"success": False, "error": "No file path provided"}))
    sys.exit(1)

fpath = sys.argv[1]
taluka_override = sys.argv[2] if len(sys.argv) > 2 else ""
original_fname = sys.argv[3] if len(sys.argv) > 3 else os.path.basename(fpath)

if not os.path.exists(fpath):
    print(json.dumps({"success": False, "error": f"File not found: {fpath}"}))
    sys.exit(1)

def clean_area(val):
    if val is None:
        return 0.0
    s = str(val).strip()
    if not s or s in ['-', '–', '—', '_', 'निरंक', 'नाही', 'None', 'nil', 'NIL']:
        return 0.0
    s_clean = re.sub(r'\s+', '', s)
    s_clean = re.sub(r'\.{2,}', '.', s_clean)
    match = re.search(r'(\d+(?:\.\d+)?)', s_clean)
    if match:
        try:
            return round(float(match.group(1)), 4)
        except Exception:
            return 0.0
    return 0.0

def clean_survey_no(val):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ['-', '–', '—', '_', 'निरंक', 'नाही', 'None', 'nil', 'NIL']:
        return None
    if s.endswith('.0'):
        s = s[:-2]
    # Check for summary phrases or headers masquerading as survey
    if any(w in s for w in [
        'सक्षम', 'बंदी', 'आदिवासी', 'हस्तांतरण', 'परवानगी', 'जमीन', 'नियम',
        'तहसिलदार', 'माहिती', 'गोषवारा', 'स्वाक्षरी', 'दिनांक', 'प्रमाणन', 'अ.क्र'
    ]):
        return None
    return s[:30]

def normalize_tenure(val, other_rights="", bandobast_tenure="", owner=""):
    s = str(val or "").strip()
    owner_str = str(owner or "").strip()
    rights_str = str(other_rights or "").strip()
    b_tenure = str(bandobast_tenure or "").strip()

    if any(k in owner_str for k in ['महाराष्ट्र शासन', 'सरकार', 'शासन', 'कालवा', 'तलाव', 'गायरान', 'सार्वजनिक', 'राखीव']):
        return "SARKAR_SHASAN"
    if 'वन' in s or 'वन' in owner_str:
        return "FOREST_JANGAL"
    if 'देवस्थान' in s or 'इनाम' in s:
        return "DEVASTHAN_INAM"
    if '2' in s or '२' in s:
        return "BHOGVATDAR_CLASS_2"
    if any(k in rights_str for k in ['आदिवासी', 'आदीवासी', '36', '३६', 'शर्तभंग', 'सक्षम प्राधिकारी']):
        return "BHOGVATDAR_CLASS_2"
    if '1' in s or '१' in s or 'भूमिधारी' in b_tenure or 'भूमिस्वामी' in b_tenure:
        return "BHOGVATDAR_CLASS_1"
    return "BHOGVATDAR_CLASS_1"

def clean_village_name(fname, sample_village):
    if sample_village and len(str(sample_village).strip()) > 2 and str(sample_village).strip() not in ['-', '_', 'None', 'गावाचे नाव', '1', '2', 'नाही']:
        clean_v = str(sample_village).strip()
        clean_v = re.sub(r'^[0-9\.\-\_\s]+', '', clean_v)
        if len(clean_v) > 2:
            return clean_v[:80]
    v = fname
    for prefix in [
        '45 मुद्द्याची माहिती', '1 ते 45 तपासणी  मुद्दे  मौजा', '1 ते 45 तपासणी मुददे',
        '1 ते 45 मुददे ची तपासणी -', '1 ते 45 मुद्दे', '1 to 45.', '1 to 46',
        '1 ते 46', '1  ते 46', '1 ते 54', '1 TO 45', '1 TO 46', '1ते45'
    ]:
        v = v.replace(prefix, '')
    v = re.sub(r'^\d+[\s_\.-]*', '', v)
    v = re.sub(r'[\s_\.-]*ok.*$', '', v, flags=re.I)
    v = re.sub(r'\.xlsx$', '', v, flags=re.I)
    v = re.sub(r'Microsoft Excel.*$', '', v, flags=re.I)
    v = v.strip()
    return v if len(v) > 1 else fname.replace('.xlsx', '')[:40]

FOOTER_STOP_KEYWORDS = [
    'माहिती तयार करणाऱ्या',
    'माहिती तपासणी करणाऱ्या',
    'स्वाक्षरी तलाठी',
    'तलाठी स्वाक्षरी',
    'स्वाक्षरी मंडळ',
    'मंडळ अधिकारी स्वाक्षरी',
    'स्वाक्षरी तहसीलदार',
    'तहसीलदार स्वाक्षरी',
    '१ ते ४५ मुद्द्यांचा गोषवारा',
    '1 ते 45 मुद्द्यांचा गोषवारा',
    'प्रमाणित करण्यात येते'
]

SUB_TYPE_NAMES = [
    "कुळ वहिवाट व शेतजमीन (विदर्भ/हैदराबाद) कायदा अन्वये कुळ घोषित केलेली जमीन",
    "इनाम व वतन जमिनी (देवस्थान जमिनी वगळुन) प्रकारातील जमिनी",
    "म.ज.म.अधि.1966 अंतर्गत विवीध योजनेअंतर्गत प्रदान/अतिक्रमण नियमानुकूल केलेल्या जमिनी (भूमिहीन, शेत मजूर, स्वातंत्र्य सैनिक इ.)",
    "म.ज.म.अधि.1966 अंतर्गत विवीध योजनेअंतर्गत प्रदान/अतिक्रमण नियमानुकूल केलेल्या जमिनी (गृहनिर्माण संस्था,औद्योगिक आस्थापना,शैक्षणिक संस्था,विशेष वसाहत,प्रकल्प इ.) प्रकारची जमिन",
    "महाराष्ट्र शेतजमिन कमाल धारणा अधिनियम 1961 अंतर्गत वाटप केलेल्या जमिनी",
    "महानगरपालिका, नगर पालिका व विविध प्राधिकरण यांच्या विकास आराखडयात समाविष्ट असलेल्या जमिनी अथवा ग्रामपंचायतकडे गुरचरण अथवा इतर प्रयोजनासाठी वर्ग केलेल्या जमिनी",
    "जमीन देवस्थान इनाम जमिनी",
    "आदिवासी खातेदारांचे मालकीच्या जमिनी",
    "महाराष्ट्र पुर्नवसन अधिनियम 1999 च्या कलम 16 अन्वये प्रदान केलेल्या जमिनी",
    "भाडेपट्टयाने दिलेली शासकीय जमिन",
    "भुदान व ग्रामदान अंतर्गत प्रदान केलेल्या जमिनी",
    "महाराष्ट्र खाजगी वने (संपादन) अधिनियम 1975 तसेच महाराष्ट्र शेतजमिन (जमिन धारणेची कमाल मर्यादा) अधिनियम 1961 अन्वये चौकशीसाठी प्रलंबित असलेल्या जमीन",
    "भुमीधारी हक्कान्वये प्राप्त झालेल्या जमिनी",
    "महाराष्ट्र शेतजमिन (जमिन धारणेची कमाल मर्यादा) अधिनियम 1961 अंतर्गत कमाल मर्यादेपेक्षा अधिक धारण करण्यास सुट दिलेल्या जमिनी",
    "वक्फ बोर्डाचे मालकीची जमीन",
    "भुसंपादन कायदयाअंतर्गत अधिग्रहित/राखीव जमीन"
]

try:
    wb = openpyxl.load_workbook(fpath, data_only=True, read_only=True)
    all_sheet_names = wb.sheetnames

    target_sheet = wb.worksheets[0]
    for s in wb.worksheets:
        s_title = s.title.strip()
        if any(k in s_title for k in ['7/12', 'तपासणी', 'मुद्दे', '1 to 45', '1 ते 45', 'Booklet', 'Master']):
            target_sheet = s
            break

    raw_rows = list(target_sheet.iter_rows(values_only=True))

    header_row_idx = 2
    for r_idx, r in enumerate(raw_rows[:8]):
        line_str = ' '.join([str(x or '').strip() for x in r[:15]])
        if any(k in line_str for k in ['सर्व्हे', 'स.नं', 'गट क्रमांक']) and any(k in line_str for k in ['क्षेत्र', 'गाव']):
            header_row_idx = r_idx
            break

    header_vals = [str(x or '').strip() for x in raw_rows[header_row_idx]]

    # 1. Landmark Column Anchoring across the header row
    col_village = next((i for i, c in enumerate(header_vals) if 'गावाचे नाव' in c or c == 'गाव'), 1)
    col_survey = next((i for i, c in enumerate(header_vals) if any(k in c for k in ['सर्व्हे क्रमांक', 'स.नं', 'गट क्रमांक', 'गट क्र'])), 2)
    col_tenure = next((i for i, c in enumerate(header_vals) if any(k in c for k in ['भोगवटदार वर्ग', 'भूधारणा', 'कब्जेदार सदरी'])), -1)
    col_owner = next((i for i, c in enumerate(header_vals) if any(k in c for k in ['भोगवटदाराचे नाव', 'खातेदाराचे नाव'])), -1)
    col_rights = next((i for i, c in enumerate(header_vals) if 'इतर अधिकार' in c), -1)

    # 2. Determine point 1 start index (Checkpoint 1 begins after Other Rights)
    col_point_start = -1
    for c_idx, cell in enumerate(header_vals):
        if c_idx > col_survey and cell == '1':
            col_point_start = c_idx
            break

    if col_point_start == -1:
        # Check row below for 'बंदोबस्त'
        if header_row_idx + 1 < len(raw_rows):
            sub_vals = [str(x or '').strip() for x in raw_rows[header_row_idx + 1]]
            for c_idx, cell in enumerate(sub_vals):
                if 'बंदोबस्त' in cell:
                    col_point_start = c_idx
                    break

    # If tenure was not explicitly found, fallback relative to survey
    if col_tenure == -1:
        col_tenure = col_survey + 2
    if col_owner == -1:
        col_owner = col_tenure + 1
    if col_rights == -1:
        col_rights = col_owner + 1
    if col_point_start == -1:
        col_point_start = col_rights + 1

    # 3. Dynamic Area Columns Identification (All columns between Survey and Tenure!)
    # No matter whether officer inserted 1, 2, 3 or 4 columns under क्षेत्र!
    area_col_indices = list(range(col_survey + 1, col_tenure))

    col_agri = -1
    col_na = -1
    col_pot = -1
    col_single_area = -1

    # Check sub-header rows (header_row_idx + 1 to header_row_idx + 2)
    for sub_r in range(header_row_idx + 1, min(header_row_idx + 3, len(raw_rows))):
        sub_row_vals = [str(x or '').strip() for x in raw_rows[sub_r]]
        for c_idx in area_col_indices:
            if c_idx < len(sub_row_vals):
                cell_text = sub_row_vals[c_idx]
                if 'अकृषक' in cell_text:
                    col_na = c_idx
                elif 'कृषक' in cell_text:
                    col_agri = c_idx
                elif 'पोटखराबा' in cell_text:
                    col_pot = c_idx

    # If sub-headers were not found but there are 2 area columns: default 1st=agri, 2nd=na
    if len(area_col_indices) == 2 and col_agri == -1 and col_na == -1:
        col_agri = area_col_indices[0]
        col_na = area_col_indices[1]
    elif len(area_col_indices) == 1:
        col_single_area = area_col_indices[0]
    elif len(area_col_indices) >= 3 and col_agri == -1:
        col_agri = area_col_indices[0]
        col_na = area_col_indices[1]
        col_pot = area_col_indices[2]

    has_dual_area = (col_agri != -1 and col_na != -1)

    # 4. Data Rows Start Index (Look for '1' or '१' in first column)
    data_start_idx = header_row_idx + 2
    for idx in range(header_row_idx + 1, min(header_row_idx + 10, len(raw_rows))):
        first_c = str(raw_rows[idx][0] or '').strip() if raw_rows[idx] else ''
        if re.match(r'^[0-9]+(\.0)?$', first_c) or first_c in ['१', '1']:
            data_start_idx = idx
            break

    # 5. Extract Sample Village
    sample_village = None
    for r in raw_rows[data_start_idx:data_start_idx + 15]:
        v_cell = str(r[col_village] or '').strip() if len(r) > col_village else ''
        if len(v_cell) > 2 and v_cell not in ['-', '_', 'None', 'गावाचे नाव', '1', '2', 'नाही']:
            sample_village = v_cell
            break

    fname = original_fname or os.path.basename(fpath)
    v_name = clean_village_name(fname, sample_village)
    if not v_name or re.match(r'^[0-9a-f]{16,}$', v_name, re.I):
        fallback = os.path.splitext(original_fname)[0] if original_fname else 'गाव'
        v_name = clean_village_name(fallback, None)

    v_stats = {
        "total712": 0,
        "inspected1950": 0,
        "govt1950": 0,
        "illegalDel": 0,
        "pokalist": 0,
        "corpAllot": 0,
        "corpShartbhang": 0,
        "privateAllot": 0,
        "privateShartbhang": 0,
        "totalGovt": 0,
        "govtAllotted": 0,
        "govtShartbhang": 0,
        "govtEncroach": 0,
        "govtMojani": 0,
        "class1Count": 0,
        "class1AreaHa": 0.0,
        "class2Count": 0,
        "class2AreaHa": 0.0,
        "govtAreaHa": 0.0,
        "shartbhangAreaHa": 0.0,
        "illegalDelAreaHa": 0.0
    }

    # Inspect other sheets for subtype associations if present (e.g. Sheet 2: DROPDOWN & 16 Sub-types)
    sheet_subtypes = {}
    if len(wb.worksheets) > 1:
        for s in wb.worksheets[1:]:
            s_title = s.title.strip()
            matched_cat = next((cat for cat in SUB_TYPE_NAMES if cat in s_title or s_title in cat), None)
            if matched_cat:
                for row_vals in s.iter_rows(values_only=True):
                    if row_vals and len(row_vals) > 1:
                        gat_cand = clean_survey_no(row_vals[1]) or clean_survey_no(row_vals[2])
                        if gat_cand:
                            sheet_subtypes[gat_cand] = matched_cat

    parcels = []
    p_offset = col_point_start

    for r_idx, r in enumerate(raw_rows[data_start_idx:], start=1):
        if not r or not any(r):
            continue

        sr_val = str(r[0] or '').strip()
        is_numeric_sr = bool(re.match(r'^[0-9]+(\.0)?$', sr_val))

        # Check footer stop ONLY on rows that are NOT numeric parcel rows,
        # and only look in the first 4 columns (where signature / summary text is placed).
        lead_cells = ' '.join([str(x or '').strip() for x in r[:4]])
        if not is_numeric_sr and any(stop_word in lead_cells for stop_word in FOOTER_STOP_KEYWORDS):
            break

        survey_raw = r[col_survey] if len(r) > col_survey else None
        owner_raw = r[col_owner] if len(r) > col_owner else None
        tenure_raw = r[col_tenure] if len(r) > col_tenure else None
        rights_raw = r[col_rights] if len(r) > col_rights else None

        # Clean Area handling (preserves raw, computes clean sum)
        agri_clean = 0.0
        na_clean = 0.0
        total_area_clean = 0.0

        if has_dual_area:
            agri_raw = r[col_agri] if len(r) > col_agri else None
            na_raw = r[col_na] if len(r) > col_na else None
            agri_clean = clean_area(agri_raw)
            na_clean = clean_area(na_raw)
            total_area_clean = round(agri_clean + (na_clean / 10000.0), 4)
        else:
            c_area_idx = col_single_area if col_single_area != -1 else (col_survey + 1)
            area_raw = r[c_area_idx] if len(r) > c_area_idx else None
            total_area_clean = clean_area(area_raw)
            agri_clean = total_area_clean

        survey_clean = clean_survey_no(survey_raw)
        if not survey_clean and not owner_raw and total_area_clean == 0.0:
            continue

        survey_final = survey_clean or f"S-{r_idx}"
        owner_clean = str(owner_raw or "अज्ञात खातेदार").strip()[:240]
        rights_clean = str(rights_raw or "नाही").strip()

        # Helper to extract exact raw string from row for points 1 to 46
        def get_raw_pt(pt_num):
            idx = p_offset + (pt_num - 1)
            if len(r) > idx and r[idx] is not None:
                val = str(r[idx]).strip()
                return val if val != 'None' else ''
            return ''

        # Store every single cell value verbatim in auditPoints!
        audit_points = {}
        for pt in range(1, 47):
            audit_points[f"point_{pt}"] = get_raw_pt(pt)

        # Backward Milestones: preserve exact raw values
        b_no = audit_points.get("point_1") or None
        b_owner = audit_points.get("point_2") or ""
        b_tenure = audit_points.get("point_3") or ""

        ek_no = audit_points.get("point_4") or None
        ek_owner = audit_points.get("point_5") or ""
        ek_tenure = audit_points.get("point_6") or ""

        res_no = audit_points.get("point_7") or None
        res_owner = audit_points.get("point_8") or ""
        res_tenure = audit_points.get("point_9") or ""

        # Checkpoints for violation flags
        col17_sarkar = audit_points.get("point_10", "")
        col18_illegal_del = audit_points.get("point_18", "")
        col19_pokalist = audit_points.get("point_19", "")
        col21_private_allot = audit_points.get("point_22", "")
        col22_encroach = audit_points.get("point_15", "")
        col23_mojani = audit_points.get("point_16", "")
        col24_corp = audit_points.get("point_23", "")
        col26_tribal = audit_points.get("point_27", "")
        col35_shartbhang = audit_points.get("point_37", "")
        col36_shartbhang_detail = audit_points.get("point_38", "")
        col_class2_to_1 = audit_points.get("point_36", "")

        tenure_final = normalize_tenure(tenure_raw, rights_clean, b_tenure, owner_clean)

        is_govt = (tenure_final == "SARKAR_SHASAN") or ('सरकार' in col17_sarkar or 'शासन' in col17_sarkar)
        is_illegal_del = col18_illegal_del in ['होय', '1', '1.0'] or ('अवैध' in col18_illegal_del)
        is_pokalist = col19_pokalist in ['होय', '1', '1.0'] or any(k in col19_pokalist for k in ['बु.ग.दे', 'ता.ग.दे', 'सं.कु'])
        is_encroach = col22_encroach in ['होय', '1', '1.0', 'आहे']
        is_mojani = col23_mojani in ['होय', '1', '1.0', 'आहे']
        is_shartbhang = (col35_shartbhang in ['होय', '1', '1.0']) or (col36_shartbhang_detail and col36_shartbhang_detail not in ['नाही', '-', '–', '—', '0', '0.0', 'निरंक', '_', 'नोंद नाही'])
        is_tribal = (col26_tribal in ['होय', '1', '1.0']) or ('आदिवासी' in rights_clean)
        is_corp_allot = (col24_corp and col24_corp not in ['नाही', '-', '–', '—', '0', '0.0', 'निरंक', '_', 'नोंद नाही'])
        is_private_allot = (col21_private_allot in ['होय', '1', '1.0']) or (col21_private_allot and col21_private_allot not in ['नाही', '-', '–', '—', '_', 'निरंक'])
        is_class2_to_1 = col_class2_to_1 in ['होय', '1', '1.0'] or ('वर्ग-१' in col_class2_to_1 or 'वर्ग 1' in col_class2_to_1)

        # Detect statutory sub-type (1-16)
        detected_subtype = sheet_subtypes.get(survey_final)
        if not detected_subtype:
            if 'कुळ' in rights_clean or audit_points.get("point_20") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[0]
            elif 'इनाम' in rights_clean or 'वतन' in rights_clean or audit_points.get("point_21") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[1]
            elif any(k in rights_clean for k in ['भूमिहीन', 'शेतमजूर', 'स्वातंत्र्य']) or audit_points.get("point_22") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[2]
            elif any(k in rights_clean for k in ['संस्था', 'उद्योग', 'शैक्षणिक', 'प्रकल्प']) or audit_points.get("point_23") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[3]
            elif any(k in rights_clean for k in ['कमाल धारणा', 'सिलिंग']) or audit_points.get("point_24") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[4]
            elif any(k in rights_clean for k in ['गुरचरण', 'चराई', 'विकास आराखडा']) or audit_points.get("point_25") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[5]
            elif 'देवस्थान' in rights_clean or tenure_final == 'DEVASTHAN_INAM' or audit_points.get("point_26") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[6]
            elif is_tribal or audit_points.get("point_27") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[7]
            elif 'पुनर्वसन' in rights_clean or audit_points.get("point_28") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[8]
            elif 'भाडेपट्टा' in rights_clean or 'लीज' in rights_clean or audit_points.get("point_29") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[9]
            elif 'भूदान' in rights_clean or 'ग्रामदान' in rights_clean or audit_points.get("point_30") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[10]
            elif 'खाजगी वन' in rights_clean or audit_points.get("point_31") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[11]
            elif 'भूमीधारी' in rights_clean or audit_points.get("point_32") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[12]
            elif 'सूट' in rights_clean or audit_points.get("point_33") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[13]
            elif 'वक्फ' in rights_clean or audit_points.get("point_34") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[14]
            elif 'भूसंपादन' in rights_clean or audit_points.get("point_35") in ['होय', '1']:
                detected_subtype = SUB_TYPE_NAMES[15]

        # Aggregate village stats
        v_stats["total712"] += 1
        v_stats["inspected1950"] += 1

        if tenure_final == 'BHOGVATDAR_CLASS_1':
            v_stats["class1Count"] += 1
            v_stats["class1AreaHa"] = round(v_stats["class1AreaHa"] + total_area_clean, 4)
        elif tenure_final == 'BHOGVATDAR_CLASS_2':
            v_stats["class2Count"] += 1
            v_stats["class2AreaHa"] = round(v_stats["class2AreaHa"] + total_area_clean, 4)

        if is_govt:
            v_stats["govt1950"] += 1
            v_stats["totalGovt"] += 1
            v_stats["govtAreaHa"] = round(v_stats["govtAreaHa"] + total_area_clean, 4)
        if is_illegal_del:
            v_stats["illegalDel"] += 1
            v_stats["illegalDelAreaHa"] = round(v_stats["illegalDelAreaHa"] + total_area_clean, 4)
        if is_pokalist:
            v_stats["pokalist"] += 1
        if is_corp_allot:
            v_stats["corpAllot"] += 1
        if is_corp_allot and is_shartbhang:
            v_stats["corpShartbhang"] += 1
        if is_private_allot:
            v_stats["privateAllot"] += 1
        if is_private_allot and is_shartbhang:
            v_stats["privateShartbhang"] += 1
        if is_govt and (is_corp_allot or is_private_allot):
            v_stats["govtAllotted"] += 1
        if is_govt and is_shartbhang:
            v_stats["govtShartbhang"] += 1
        if is_govt and is_encroach:
            v_stats["govtEncroach"] += 1
        if is_govt and is_mojani:
            v_stats["govtMojani"] += 1
        if is_shartbhang:
            v_stats["shartbhangAreaHa"] = round(v_stats["shartbhangAreaHa"] + total_area_clean, 4)

        safe_survey = re.sub(r'[^0-9]', '', survey_final) or f"{r_idx}"
        taluka_prefix = (taluka_override or "DIS")[:3].upper()
        upi = f"MH-CHA-{taluka_prefix}-UPL-{int(safe_survey[:4]) if safe_survey.isdigit() else 1:04d}-{r_idx:02d}"

        parcels.append({
            "upi": upi,
            "gatNumber": survey_final,
            "oldSurveyNo": b_no,
            "agriAreaHa": agri_clean,
            "naAreaSqM": na_clean,
            "totalAreaHa": total_area_clean,
            "tenureClass": tenure_final,
            "ownerName": owner_clean,
            "otherRights": rights_clean,
            "subTypeCategory": detected_subtype,
            "auditPoints": audit_points,
            "backward": {
                "settlement1950": {
                    "no": b_no,
                    "owner": b_owner,
                    "tenure": b_tenure
                },
                "ekatrikaran1972": {
                    "no": ek_no,
                    "owner": ek_owner,
                    "tenure": ek_tenure
                },
                "resurvey1995": {
                    "no": res_no,
                    "owner": res_owner,
                    "tenure": res_tenure
                }
            },
            "flags": {
                "isGovt": is_govt,
                "isIllegalDeletion": is_illegal_del,
                "isPokalist": is_pokalist,
                "isEncroached": is_encroach,
                "isMojaniDone": is_mojani,
                "isShartbhang": is_shartbhang,
                "isTribal": is_tribal,
                "isCorpAllot": is_corp_allot,
                "isPrivateAllot": is_private_allot,
                "isClass2ToClass1": is_class2_to_1
            }
        })

    unique_gats = list(dict.fromkeys([p['gatNumber'] for p in parcels if p.get('gatNumber')]))

    output = {
        "success": True,
        "villageName": v_name,
        "taluka": taluka_override,
        "sheetCount": len(all_sheet_names),
        "sheetNames": all_sheet_names,
        "detectedLayout": "dual_area" if has_dual_area else ("multi_area" if len(area_col_indices) > 2 else "single_area"),
        "areaColumnsCount": len(area_col_indices),
        "stats": v_stats,
        "parcels": parcels,
        "uniqueGatNumbers": unique_gats
    }
    print(json.dumps(output, ensure_ascii=False))

except Exception as e:
    print(json.dumps({"success": False, "error": str(e)}))
