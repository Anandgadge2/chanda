import openpyxl
import os
import glob
import sys
import json
import re
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

data_dir = r"d:\chanda\data\gondpipri_real_data"
excel_files = sorted(glob.glob(os.path.join(data_dir, "*.xlsx")))

def clean_area(val):
    if val is None:
        return 0.0
    s = str(val).strip()
    if not s or s in ['-', '_', 'निरंक']:
        return 0.0
    s_clean = re.sub(r'\s+', '', s)
    s_clean = re.sub(r'\.{2,}', '.', s_clean)
    match = re.search(r'(\d+(?:\.\d+)?)', s_clean)
    if match:
        try:
            return round(float(match.group(1)), 4)
        except:
            return 0.0
    return 0.0

def clean_survey_no(val):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ['-', '_', 'निरंक']:
        return None
    if s.endswith('.0'):
        s = s[:-2]
    # Check if text is a remark instead of survey no
    if len(s) > 30 or any(w in s for w in ['सक्षम', 'बंदी', 'आदिवासी', 'हस्तांतरण', 'परवानगी', 'जमीन', 'नियम']):
        return None
    return s[:30]

def normalize_tenure(val, other_rights, bandobast_tenure, owner):
    s = str(val or "").strip()
    owner_str = str(owner or "").strip()
    rights_str = str(other_rights or "").strip()
    b_tenure = str(bandobast_tenure or "").strip()

    if any(k in owner_str for k in ['महाराष्ट्र शासन', 'सरकार', 'शासन', 'कालवा', 'तलाव', 'गायरान']):
        return "SARKAR_SHASAN"
    if 'वन' in s or 'वन' in owner_str:
        return "FOREST_JANGAL"
    if 'देवस्थान' in s or 'इनाम' in s:
        return "DEVASTHAN_INAM"
    if '2' in s or '२' in s:
        return "BHOGVATDAR_CLASS_2"
    if any(k in rights_str for k in ['आदिवासी', 'आदीवासी', '36', '३६', 'शर्तभंग', 'सक्षम प्राधिकारी']):
        return "BHOGVATDAR_CLASS_2"
    if '1' in s or '१' in s or 'भूमिधारी' in b_tenure:
        return "BHOGVATDAR_CLASS_1"
    
    return "BHOGVATDAR_CLASS_1"

parcels_output = []
village_rollups = defaultdict(lambda: defaultdict(int))

for f_idx, fpath in enumerate(excel_files, 1):
    fname = os.path.basename(fpath)
    wb = openpyxl.load_workbook(fpath, data_only=True)
    ws = wb[wb.sheetnames[0]]
    
    v_code = f"GON-{f_idx:03d}"
    raw_vname = fname.split('_')[1] if '_' in fname else fname
    raw_vname = raw_vname.replace('.xlsx', '').strip()
    current_village = raw_vname
    
    saza_name = "साजा"
    if 'सज' in fname or 'साजा' in fname:
        parts = fname.split('सज') if 'सज' in fname else fname.split('साजा')
        if len(parts) > 1:
            saza_name = ("साजा " + parts[1].replace('.xlsx', '').replace('-', '').strip())[:60]

    row_count = 0
    for r in range(7, ws.max_row + 1):
        v_cell = ws.cell(r, 2).value
        survey_raw = ws.cell(r, 3).value
        area_raw = ws.cell(r, 4).value
        tenure_raw = ws.cell(r, 5).value
        owner_raw = ws.cell(r, 6).value
        rights_raw = ws.cell(r, 7).value
        
        if not any([survey_raw, area_raw, owner_raw]):
            continue
            
        if v_cell and str(v_cell).strip() and str(v_cell).strip() not in ['_', '-']:
            current_village = str(v_cell).strip()[:100]
            
        row_count += 1
        survey_clean = clean_survey_no(survey_raw) or f"S-{row_count}"
        area_clean = clean_area(area_raw)
        owner_clean = str(owner_raw or "अज्ञात खातेदार").strip()[:240]
        rights_clean = str(rights_raw or "नाही").strip()
        
        # 1950 Bandobast (Cols 8-10)
        b_cell = ws.cell(r, 8).value
        b_no = clean_survey_no(b_cell)
        # if b_cell had text remark, append to rights
        if b_cell and not b_no and len(str(b_cell).strip()) > 5:
            rights_clean = (rights_clean + " | " + str(b_cell).strip()).strip()

        b_owner = str(ws.cell(r, 9).value or "").strip()[:240]
        b_tenure = str(ws.cell(r, 10).value or "").strip()[:60]
        
        # Consolidation (Cols 11-13)
        c_no = clean_survey_no(ws.cell(r, 11).value)
        c_owner = str(ws.cell(r, 12).value or "").strip()[:240]
        c_tenure = str(ws.cell(r, 13).value or "").strip()[:60]
        
        # Resurvey (Cols 14-16)
        r_no = clean_survey_no(ws.cell(r, 14).value)
        r_owner = str(ws.cell(r, 15).value or "").strip()[:240]
        r_tenure = str(ws.cell(r, 16).value or "").strip()[:60]
        
        # Audit points (Cols 17 to 52)
        audit_points = {}
        for c in range(17, min(ws.max_column + 1, 53)):
            p_val = str(ws.cell(r, c).value or "").strip()
            audit_points[f"point_{c-7}"] = p_val
            
        tenure_final = normalize_tenure(tenure_raw, rights_clean, b_tenure, owner_clean)
        
        col17 = str(ws.cell(r, 17).value or "").strip()
        col18_illegal_del = str(ws.cell(r, 25).value or "").strip()
        col19_pokalist = str(ws.cell(r, 26).value or "").strip()
        col21_private_allot = str(ws.cell(r, 21).value or "").strip()
        col22_encroach = str(ws.cell(r, 22).value or "").strip()
        col23_mojani = str(ws.cell(r, 23).value or "").strip()
        col24_corporation = str(ws.cell(r, 31).value or "").strip()
        col26_tribal = str(ws.cell(r, 33).value or "").strip()
        col35_shartbhang_class = str(ws.cell(r, 42).value or "").strip()
        col36_shartbhang_use = str(ws.cell(r, 43).value or "").strip()
        
        is_govt = (tenure_final == "SARKAR_SHASAN") or ('सरकार' in col17 or 'शासन' in col17)
        is_illegal_del = col18_illegal_del in ['होय', '1', '1.0'] or ('अवैध' in col18_illegal_del)
        is_pokalist = col19_pokalist in ['होय', '1', '1.0'] or ('बु.ग.दे' in col19_pokalist or 'ता.ग.दे' in col19_pokalist or 'सं.कु' in col19_pokalist)
        is_encroach = col22_encroach in ['होय', '1', '1.0']
        is_mojani = col23_mojani in ['होय', '1', '1.0'] or ('आहे' in col23_mojani)
        is_shartbhang = (col35_shartbhang_class in ['होय', '1', '1.0']) or (col36_shartbhang_use and col36_shartbhang_use not in ['नाही', '-', '0', '0.0', 'निरंक', '_', 'नोंद नाही'])
        is_tribal = (col26_tribal in ['होय', '1', '1.0']) or ('आदिवासी' in rights_clean) or ('आदीवासी' in rights_clean)
        is_corp_allot = (col24_corporation and col24_corporation not in ['नाही', '-', '0', '0.0', 'निरंक', '_', 'नोंद नाही'])
        is_private_allot = (col21_private_allot in ['होय', '1', '1.0']) or (col21_private_allot and col21_private_allot not in ['नाही', '-', '_', 'निरंक'])

        safe_survey = re.sub(r'[^0-9]', '', survey_clean) or "0001"
        upi = f"MH-CHA-GON-{f_idx:03d}-{int(safe_survey[:4]):04d}-{row_count:02d}"
        
        parcel_obj = {
            "upi": upi,
            "district": "Chandrapur",
            "taluka": "Gondpipri",
            "revenueCircle": saza_name,
            "villageCode": v_code,
            "villageName": current_village,
            "oldSurveyNo": b_no,
            "gatNumber": survey_clean,
            "totalAreaHa": area_clean,
            "tenureClass": tenure_final,
            "ownerName": owner_clean,
            "otherRights": rights_clean,
            "backward": {
                "bandobast": {"no": b_no, "owner": b_owner, "tenure": b_tenure},
                "consolidation": {"no": c_no, "owner": c_owner, "tenure": c_tenure},
                "resurvey": {"no": r_no, "owner": r_owner, "tenure": r_tenure}
            },
            "auditPoints": audit_points,
            "flags": {
                "isGovt": is_govt,
                "isIllegalDeletion": is_illegal_del,
                "isPokalist": is_pokalist,
                "isEncroached": is_encroach,
                "isMojaniDone": is_mojani,
                "isShartbhang": is_shartbhang,
                "isTribal": is_tribal,
                "isCorpAllot": is_corp_allot,
                "isPrivateAllot": is_private_allot
            }
        }
        parcels_output.append(parcel_obj)
        
        # Rollups updating
        v_stats = village_rollups[current_village]
        v_stats["total712"] += 1
        v_stats["inspected1950"] += 1
        if is_govt: v_stats["govt1950"] += 1
        if is_illegal_del: v_stats["illegalDel"] += 1
        if is_pokalist: v_stats["pokalist"] += 1
        if is_corp_allot: v_stats["corpAllot"] += 1
        if is_corp_allot and is_shartbhang: v_stats["corpShartbhang"] += 1
        if is_private_allot: v_stats["privateAllot"] += 1
        if is_private_allot and is_shartbhang: v_stats["privateShartbhang"] += 1
        if is_govt: v_stats["totalGovt"] += 1
        if is_govt and (is_corp_allot or is_private_allot): v_stats["govtAllotted"] += 1
        if is_govt and is_shartbhang: v_stats["govtShartbhang"] += 1
        if is_govt and is_encroach: v_stats["govtEncroach"] += 1
        if is_govt and is_mojani: v_stats["govtMojani"] += 1

print(f"Processed {len(parcels_output)} parcels across {len(village_rollups)} villages.")

out_data = {
    "totalParcels": len(parcels_output),
    "parcels": parcels_output,
    "villageRollups": village_rollups
}

out_file = r"d:\chanda\data\gondpipri_cleaned_payload.json"
with open(out_file, "w", encoding="utf-8") as f:
    json.dump(out_data, f, ensure_ascii=False, indent=2)

print(f"Exported clean payload to {out_file} ({os.path.getsize(out_file):,} bytes)")
