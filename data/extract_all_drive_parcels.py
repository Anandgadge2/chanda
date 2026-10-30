import os
import sys
import re
import json
import time
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

TALUKA_FOLDERS = [
    (r"d:\chanda\data\all_talukas_raw\Ballarpur", "Ballarpur"),
    (r"d:\chanda\data\all_talukas_raw\Bhadravati", "Bhadravati"),
    (r"d:\chanda\data\all_talukas_raw\Bramhapuri", "Brahmapuri"),
    (r"d:\chanda\data\all_talukas_raw\Chandrapur", "Chandrapur"),
    (r"d:\chanda\data\all_talukas_raw\Chimur", "Chimur"),
    (r"d:\chanda\data\all_talukas_raw\Jivati", "Jivati"),
    (r"d:\chanda\data\all_talukas_raw\Korpana", "Korpana"),
    (r"d:\chanda\data\all_talukas_raw\Mul", "Mul"),
    (r"d:\chanda\data\all_talukas_raw\Nagbhid", "Nagbhid"),
    (r"d:\chanda\data\all_talukas_raw\Pombhurna", "Pombhurna"),
    (r"d:\chanda\data\all_talukas_raw\Rajura", "Rajura"),
    (r"d:\chanda\data\all_talukas_raw\Saoli", "Saoli"),
    (r"d:\chanda\data\all_talukas_raw\Sindewahi", "Sindewahi"),
    (r"d:\chanda\data\all_talukas_raw\Warora", "Warora"),
    (r"d:\chanda\data\gondpipri_real_data", "Gondpipri")
]

FOOTER_STOP_KEYWORDS = [
    'माहिती तयार करणाऱ्या',
    'माहिती तपासणी करणाऱ्या',
    'स्वाक्षरी तलाठी',
    'स्वाक्षरी मंडळ अधिकारी',
    'तहसिलदार',
    'नायब तहसिलदार',
    'एकूण बेरीज',
    'गोषवारा'
]

def clean_area(val):
    if val is None:
        return 0.0
    s = str(val).strip().replace(' ', '').replace('..', '.')
    if not s or s in ['-', '–', '—', '_', 'निरंक', 'नाही', 'None', 'nil', 'NIL', 'null']:
        return 0.0
    # Handle formats like "0.66.00" -> "0.66"
    parts = s.split('.')
    if len(parts) > 2:
        s = parts[0] + '.' + parts[1]
    m = re.search(r'(\d+(?:\.\d+)?)', s)
    if m:
        try:
            return round(float(m.group(1)), 4)
        except:
            return 0.0
    return 0.0

def clean_survey_no(val):
    if val is None:
        return None
    s = str(val).strip()
    if not s or s in ['-', '–', '—', '_', 'निरंक', 'नाही', 'None', 'nil', 'NIL', 'null']:
        return None
    if s.endswith('.0'):
        s = s[:-2]
    if s.endswith('.') and len(s) > 1:
        s = s[:-1]
    invalid_words = [
        'सक्षम', 'बंदी', 'आदिवासी', 'हस्तांतरण', 'परवानगी', 'जमीन', 'नियम',
        'तहसिलदार', 'माहिती', 'गोषवारा', 'स्वाक्षरी', 'दिनांक', 'प्रमाणन', 'अ.क्र'
    ]
    if any(w in s for w in invalid_words):
        return None
    return s[:32]

def clean_village_name(fname, sample_village):
    if sample_village and len(str(sample_village).strip()) > 2:
        cv = str(sample_village).strip()
        cv = re.sub(r'^[0-9.\-_ \t]+', '', cv)
        if len(cv) > 2 and cv not in ['-', '_', 'None', 'गावाचे नाव', '1', '2', 'नाही']:
            return cv[:80]
    v = fname or 'गाव'
    prefixes = [
        '45 मुद्द्याची माहिती', '1 ते 45 तपासणी  मुद्दे  मौजा', '1 ते 45 तपासणी मुददे',
        '1 ते 45 मुददे ची तपासणी -', '1 ते 45 मुद्दे', '1 to 45.', '1 to 46',
        '1 ते 46', '1  ते 46', '1 ते 54', '1 TO 45', '1 TO 46', '1ते45'
    ]
    for p in prefixes:
        v = v.replace(p, '')
    v = re.sub(r'^\d+[\s_.-]*', '', v)
    v = re.sub(r'[\s_.-]*ok.*$', '', v, flags=re.I)
    v = re.sub(r'\.xlsx$', '', v, flags=re.I)
    v = v.strip()
    return v[:40] if len(v) > 1 else 'गाव'

def normalize_tenure(val, rights, bandobast_tenure, owner):
    s = str(val or '').strip()
    owner_str = str(owner or '').strip()
    rights_str = str(rights or '').strip()
    b_tenure = str(bandobast_tenure or '').strip()

    if any(k in owner_str for k in ['महाराष्ट्र शासन', 'सरकार', 'शासन', 'कालवा', 'तलाव', 'गायरान', 'सार्वजनिक', 'राखीव']):
        return 'SARKAR_SHASAN'
    if 'वन' in s or 'वन' in owner_str:
        return 'FOREST_JANGAL'
    if any(k in s for k in ['देवस्थान', 'इनाम']):
        return 'DEVASTHAN_INAM'
    if '2' in s or '२' in s:
        return 'BHOGVATDAR_CLASS_2'
    if any(k in rights_str for k in ['आदिवासी', 'आदीवासी', '36', '३६', 'शर्तभंग', 'सक्षम प्राधिकारी']):
        return 'BHOGVATDAR_CLASS_2'
    if '1' in s or '१' in s or 'भूमिधारी' in b_tenure or 'भूमिस्वामी' in b_tenure:
        return 'BHOGVATDAR_CLASS_1'
    return 'BHOGVATDAR_CLASS_1'

def process_single_file(fp, taluka_name, file_idx):
    wb = openpyxl.load_workbook(fp, read_only=True, data_only=True)
    target_sheet = None
    
    # Priority 1: sheet name containing relevant keywords
    for s_name in wb.sheetnames:
        if any(k in s_name for k in ['7/12', 'तपासणी', 'मुद्दे', '1 to 45', '1 ते 45', 'Master', 'Sheet1', 'saywan']):
            ws = wb[s_name]
            first_rows = list(ws.iter_rows(values_only=True, max_row=5))
            if first_rows and any(any(r) for r in first_rows):
                target_sheet = ws
                break
                
    if not target_sheet:
        target_sheet = wb.active

    rows = list(target_sheet.iter_rows(values_only=True))
    wb.close()
    
    if not rows or len(rows) < 2:
        return None

    # Header Row detection
    header_row_idx = -1
    for r_idx in range(min(8, len(rows))):
        line_str = ' '.join(str(c or '').strip() for c in rows[r_idx][:15])
        if any(k in line_str for k in ['सर्व्हे', 'स.नं', 'गट क्रमांक', 'अ.क्र']) and any(k in line_str for k in ['क्षेत्र', 'गाव', 'भोगवटदार']):
            header_row_idx = r_idx
            break
            
    if header_row_idx == -1:
        header_row_idx = 2 if len(rows) > 2 else 0

    h_row = [str(c or '').strip() for c in rows[header_row_idx]]

    # Landmark columns
    col_village = -1
    col_survey = -1
    col_tenure = -1
    col_owner = -1
    col_rights = -1

    for i, c in enumerate(h_row):
        if any(k in c for k in ['गावाचे नाव', 'गाव']) and col_village == -1:
            col_village = i
        elif any(k in c for k in ['सर्व्हे क्रमांक', 'स.नं', 'गट क्रमांक', 'गट क्र']) and col_survey == -1:
            col_survey = i
        elif any(k in c for k in ['भोगवटदार वर्ग', 'भूधारणा', 'कब्जेदार सदरी']) and col_tenure == -1:
            col_tenure = i
        elif any(k in c for k in ['भोगवटदाराचे नाव', 'खातेदाराचे नाव']) and col_owner == -1:
            col_owner = i
        elif 'इतर अधिकार' in c and col_rights == -1:
            col_rights = i

    if col_village == -1: col_village = 1
    if col_survey == -1: col_survey = 2
    if col_tenure == -1: col_tenure = col_survey + 2
    if col_owner == -1: col_owner = col_tenure + 1
    if col_rights == -1: col_rights = col_owner + 1

    # Detect col_point1
    col_point1 = -1
    for r_idx in range(min(header_row_idx + 4, len(rows))):
        r = rows[r_idx]
        for c_idx in range(len(r)):
            s = str(r[c_idx] or '').strip()
            if re.match(r'^1(\.0)?$', s):
                next_val = str(r[c_idx + 1] or '').strip() if c_idx + 1 < len(r) else ''
                if re.match(r'^2(\.0)?$', next_val):
                    col_point1 = c_idx
                    break
        if col_point1 != -1:
            break

    if col_point1 == -1:
        for r_idx in range(min(header_row_idx + 4, len(rows))):
            r = rows[r_idx]
            for c_idx in range(len(r)):
                if 'बंदोबस्त' in str(r[c_idx] or ''):
                    col_point1 = c_idx
                    break
            if col_point1 != -1:
                break

    if col_point1 == -1:
        col_point1 = col_rights + 1

    # Detect area columns
    col_agri = -1
    col_na = -1
    col_single_area = -1

    for sub_r in range(header_row_idx, min(header_row_idx + 4, len(rows))):
        sub_row = [str(c or '').strip() for c in rows[sub_r]]
        for c_idx in range(col_survey + 1, min(col_tenure, len(sub_row))):
            ct = sub_row[c_idx]
            if 'अकृषक' in ct or 'अकृषक' in ct:
                col_na = c_idx
            elif 'कृषक' in ct:
                col_agri = c_idx

    if col_agri == -1 and col_na == -1:
        col_single_area = col_survey + 1

    # Determine data start row
    data_start = header_row_idx + 1
    while data_start < len(rows):
        first_cell = str(rows[data_start][0] or '').strip()
        second_cell = str(rows[data_start][1] or '').strip() if len(rows[data_start]) > 1 else ''
        # If starts with number (1, 2, 3) or has valid survey number
        if re.match(r'^\d+(\.0)?$', first_cell) or clean_survey_no(rows[data_start][col_survey] if len(rows[data_start]) > col_survey else None):
            break
        data_start += 1

    sample_village_name = None
    for r in rows[data_start:data_start + 15]:
        if len(r) > col_village and r[col_village]:
            cv = str(r[col_village]).strip()
            if len(cv) > 2 and cv not in ['-', '_', 'None', 'गावाचे नाव', '1', '2', 'नाही']:
                sample_village_name = cv
                break

    fname = os.path.basename(fp)
    final_village_name = clean_village_name(fname, sample_village_name)
    taluka_prefix = taluka_name[:3].upper()
    village_code = f"{taluka_prefix}-{file_idx:03d}"

    parcels = []
    
    for r_idx in range(data_start, len(rows)):
        r = rows[r_idx]
        if not r or not any(r):
            continue

        lead_str = ' '.join(str(c or '').strip() for c in r[:4])
        sr_val = str(r[0] or '').strip()
        if not re.match(r'^\d+(\.0)?$', sr_val) and any(kw in lead_str for kw in FOOTER_STOP_KEYWORDS):
            break

        survey_raw = r[col_survey] if len(r) > col_survey else None
        owner_raw = r[col_owner] if len(r) > col_owner else None
        tenure_raw = r[col_tenure] if len(r) > col_tenure else None
        rights_raw = r[col_rights] if len(r) > col_rights else None

        agri_clean = 0.0
        na_clean = 0.0
        total_area = 0.0

        if col_agri != -1 or col_na != -1:
            if col_agri != -1 and len(r) > col_agri:
                agri_clean = clean_area(r[col_agri])
            if col_na != -1 and len(r) > col_na:
                na_clean = clean_area(r[col_na])
            total_area = round(agri_clean + (na_clean / 10000.0), 4)
        else:
            if col_single_area != -1 and len(r) > col_single_area:
                total_area = clean_area(r[col_single_area])
                agri_clean = total_area

        survey_clean = clean_survey_no(survey_raw)
        if not survey_clean and not owner_raw and total_area == 0.0:
            continue

        survey_final = survey_clean or f"S-{len(parcels) + 1}"
        owner_clean = str(owner_raw or "अज्ञात खातेदार").strip()[:240]
        rights_clean = str(rights_raw or "नाही").strip()

        # Audit points 1 to 46
        audit_points = {}
        for pt in range(1, 47):
            idx = col_point1 + (pt - 1)
            val = ''
            if len(r) > idx and r[idx] is not None:
                s_val = str(r[idx]).strip()
                if s_val not in ['None', 'null']:
                    val = s_val
            audit_points[f"point_{pt}"] = val

        # Milestone 1950 Bandobast
        b_no = audit_points.get('point_1') or None
        b_owner = audit_points.get('point_2') or ''
        b_tenure = audit_points.get('point_3') or ''

        # Milestone 1972 Ekatrikaran (Consolidation)
        ek_no = audit_points.get('point_4') or None
        ek_owner = audit_points.get('point_5') or ''
        ek_tenure = audit_points.get('point_6') or ''

        # Milestone 1995 Resurvey (Punarmajani)
        res_no = audit_points.get('point_7') or None
        res_owner = audit_points.get('point_8') or ''
        res_tenure = audit_points.get('point_9') or ''

        # Flags detection
        col17_sarkar = audit_points.get('point_10') or ''
        col18_illegal_del = audit_points.get('point_18') or ''
        col19_pokalist = audit_points.get('point_19') or ''
        col21_private_allot = audit_points.get('point_22') or ''
        col22_encroach = audit_points.get('point_15') or ''
        col23_mojani = audit_points.get('point_16') or ''
        col35_class2to1 = audit_points.get('point_36') or ''
        col36_shartbhang = audit_points.get('point_37') or ''
        col37_tribal1 = audit_points.get('point_38') or ''
        col38_tribal2 = audit_points.get('point_39') or ''
        col_tribal_owner = audit_points.get('point_27') or ''

        is_govt = any(k in col17_sarkar or k in owner_clean for k in ['शासकीय', 'सरकार', 'महाराष्ट्र शासन', 'गायरान', 'पडीत', 'गावठान'])
        is_illegal_del = any(k in col18_illegal_del for k in ['होय', 'आहे', 'तफावत', 'कमी'])
        is_pokalist = any(k in col19_pokalist or k in rights_clean for k in ['होय', 'आहे', 'नोंद', 'बु.ग.दे', 'ता.ग.दे'])
        is_encroach = any(k in col22_encroach or k in rights_clean for k in ['होय', 'आहे', 'अतिक्रमण'])
        is_mojani = any(k in col23_mojani for k in ['होय', 'आहे', 'झाली'])
        is_shartbhang = any(k in col36_shartbhang or k in rights_clean for k in ['होय', 'आहे', 'उल्लंघन', 'शर्तभंग'])
        is_tribal = (any(k in col37_tribal1 or k in col38_tribal2 or k in col_tribal_owner for k in ['होय', 'आहे']) or
                     any(k in rights_clean or k in owner_clean for k in ['आदिवासी', 'आदीवासी', 'कलम ३६', 'कलम 36', '36A', '३६अ']))
        is_private_allot = any(k in col21_private_allot for k in ['होय', 'आहे'])
        is_class2to1 = any(k in col35_class2to1 or k in rights_clean for k in ['होय', 'आहे', 'रूपांतर', 'रुपांतरीत'])

        tenure_final = normalize_tenure(tenure_raw, rights_clean, b_tenure, owner_clean)
        if is_govt:
            tenure_final = 'SARKAR_SHASAN'

        safe_gat = re.sub(r'[^0-9a-zA-Z]', '-', str(survey_final))[:16]
        upi = f"MH-CHA-{taluka_prefix}-{village_code}-{safe_gat}-{len(parcels) + 1:04d}"[:64]

        parcels.append({
            'upi': upi,
            'district': 'Chandrapur',
            'taluka': taluka_name,
            'revenueCircle': f"मंडळ {final_village_name[:20]}",
            'villageCode': village_code,
            'villageName': final_village_name,
            'gatNumber': survey_final,
            'hissaNumber': None,
            'oldSurveyNo': b_no,
            'totalAreaHa': total_area,
            'agriAreaHa': agri_clean,
            'naAreaSqM': na_clean,
            'tenureClass': tenure_final,
            'ownerName': owner_clean,
            'otherRights': rights_clean,
            'subTypeCategory': None,
            'auditPoints': audit_points,
            'backward': {
                'settlement1950': {
                    'no': b_no,
                    'owner': b_owner,
                    'tenure': b_tenure
                },
                'ekatrikaran1972': {
                    'no': ek_no,
                    'owner': ek_owner,
                    'tenure': ek_tenure
                },
                'resurvey1995': {
                    'no': res_no,
                    'owner': res_owner,
                    'tenure': res_tenure
                }
            },
            'flags': {
                'isGovt': is_govt,
                'isIllegalDeletion': is_illegal_del,
                'isPokalist': is_pokalist,
                'isEncroached': is_encroach,
                'isMojaniDone': is_mojani,
                'isShartbhang': is_shartbhang,
                'isTribal': is_tribal,
                'isCorpAllot': False,
                'isPrivateAllot': is_private_allot,
                'isClass2ToClass1': is_class2to1
            }
        })

    return {
        'villageName': final_village_name,
        'taluka': taluka_name,
        'villageCode': village_code,
        'file': fname,
        'parcels': parcels
    }

def main():
    out_dir = r"d:\chanda\data\seeded_talukas"
    os.makedirs(out_dir, exist_ok=True)
    
    t_start = time.time()
    total_parcels_all = 0
    total_villages_all = 0
    taluka_summaries = {}

    for folder, t_name in TALUKA_FOLDERS:
        if not os.path.exists(folder):
            print(f"Skipping {folder} (does not exist)")
            continue

        print(f"\n==========================================")
        print(f"Extracting Taluka: {t_name}")
        print(f"==========================================")

        files = [f for f in os.listdir(folder) if f.endswith('.xlsx') and not f.startswith('~$')]
        taluka_parcels = []
        taluka_villages = []

        for f_idx, fname in enumerate(files, 1):
            fp = os.path.join(folder, fname)
            try:
                res = process_single_file(fp, t_name, f_idx)
                if res and res['parcels']:
                    taluka_parcels.extend(res['parcels'])
                    taluka_villages.append({
                        'villageName': res['villageName'],
                        'villageCode': res['villageCode'],
                        'file': fname,
                        'parcelCount': len(res['parcels'])
                    })
                    print(f"  [{f_idx}/{len(files)}] {res['villageName']} ({fname}): {len(res['parcels'])} parcels")
            except Exception as e:
                print(f"  [ERROR] {fname}: {e}")

        # Save taluka parcels JSON
        out_path = os.path.join(out_dir, f"{t_name}.json")
        with open(out_path, 'w', encoding='utf-8') as out_f:
            json.dump(taluka_parcels, out_f, ensure_ascii=False, indent=2)

        total_parcels_all += len(taluka_parcels)
        total_villages_all += len(taluka_villages)
        taluka_summaries[t_name] = {
            'villageCount': len(taluka_villages),
            'parcelCount': len(taluka_parcels),
            'villages': taluka_villages
        }
        print(f"Saved {len(taluka_parcels)} parcels for {t_name} to {out_path}")

    # Save summary
    summary_path = os.path.join(out_dir, "all_talukas_summary.json")
    with open(summary_path, 'w', encoding='utf-8') as sf:
        json.dump({
            'totalVillages': total_villages_all,
            'totalParcels': total_parcels_all,
            'durationSeconds': round(time.time() - t_start, 2),
            'talukas': taluka_summaries
        }, sf, ensure_ascii=False, indent=2)

    t_end = time.time()
    print(f"\n==========================================")
    print(f"Extraction Completed in {t_end - t_start:.2f} seconds!")
    print(f"Total Villages Extracted: {total_villages_all}")
    print(f"Total Parcels Extracted: {total_parcels_all}")
    print(f"Summary saved to: {summary_path}")
    print(f"==========================================")

if __name__ == '__main__':
    main()
